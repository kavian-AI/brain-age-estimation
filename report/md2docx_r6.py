#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Markdown -> right-to-left Persian Word document (with image support)."""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Inches

SRC = r"C:\Users\Sibetalaee\Desktop\PhD\Brain Age Estimation\report\Brain_Age_6.md"
OUT = r"C:\Users\Sibetalaee\Desktop\PhD\Brain Age Estimation\report\Brain_Age_6.docx"

FA_FONT = "Tahoma"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x5F, 0x6D)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def para_rtl(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        jc = pPr.find(qn("w:jc"))
        (jc.addprevious(bidi) if jc is not None else pPr.append(bidi))


def run_rtl(r):
    rPr = r._element.get_or_add_rPr()
    if rPr.find(qn("w:rtl")) is None:
        rPr.append(OxmlElement("w:rtl"))


def run_font(r, latin=FA_FONT, cs=FA_FONT, size=None):
    r.font.name = latin
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:cs"), cs)
    if size is not None:
        r.font.size = Pt(size)
        rPr.append(_el("w:szCs", val=str(int(size * 2))))


def bold_cs(r):
    r.bold = True
    r._element.get_or_add_rPr().append(OxmlElement("w:bCs"))


def table_rtl(t):
    tblPr = t._tbl.tblPr
    if tblPr.find(qn("w:bidiVisual")) is None:
        tblPr.append(OxmlElement("w:bidiVisual"))


def shade(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(
        _el("w:shd", val="clear", color="auto", fill=hexcolor))


def cell_borders(cell, color="BFC7D1"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        b.append(_el("w:" + e, val="single", sz="4", space="0", color=color))
    tcPr.append(b)


def para_border_bottom(p, color="1F3A5F", sz="12"):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    b.append(_el("w:bottom", val="single", sz=sz, space="4", color=color))
    pPr.append(b)


def para_shading(p, hexcolor):
    p._p.get_or_add_pPr().append(_el("w:shd", val="clear", color="auto", fill=hexcolor))


def keep_with_next(p):
    p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def add_field(p, instr, placeholder=""):
    p.add_run()._element.append(_el("w:fldChar", fldCharType="begin"))
    r2 = p.add_run()._element
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2.append(it)
    p.add_run()._element.append(_el("w:fldChar", fldCharType="separate"))
    if placeholder:
        r = p.add_run(placeholder)
        run_font(r, size=10)
        # Left as an LTR run on purpose: an RTL run makes Word shape the page
        # number with Arabic-Indic glyphs (١١), which clashes with the Persian
        # digits (۱۱) used throughout the body.
        r._element.get_or_add_rPr().append(_el("w:rtl", val="0"))
    p.add_run()._element.append(_el("w:fldChar", fldCharType="end"))


def add_page_numbers(doc):
    p = doc.sections[0].footer.paragraphs[0]
    para_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", "1")


def set_doc_defaults(doc):
    dd = doc.styles.element.find(qn("w:docDefaults"))
    rd = dd.find(qn("w:rPrDefault"))
    rPr = rd.find(qn("w:rPr")) or OxmlElement("w:rPr")
    if rd.find(qn("w:rPr")) is None:
        rd.append(rPr)
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), FA_FONT)
    rf.set(qn("w:hAnsi"), FA_FONT)
    rf.set(qn("w:cs"), FA_FONT)
    rPr.insert(0, rf)
    rPr.append(_el("w:sz", val="21"))
    rPr.append(_el("w:szCs", val="21"))
    rPr.append(OxmlElement("w:rtl"))

    pd = dd.find(qn("w:pPrDefault"))
    pPr = pd.find(qn("w:pPr")) or OxmlElement("w:pPr")
    if pd.find(qn("w:pPr")) is None:
        pd.append(pPr)
    pPr.append(OxmlElement("w:bidi"))

    doc.sections[0]._sectPr.append(OxmlElement("w:bidi"))

    n = doc.styles["Normal"]
    n.font.size = Pt(10.5)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.3

    for lvl, size in ((1, 16), (2, 13), (3, 11.5), (4, 11)):
        st = doc.styles["Heading %d" % lvl]
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.font.bold = True
        rp = st.element.get_or_add_rPr()
        f = OxmlElement("w:rFonts")
        f.set(qn("w:ascii"), FA_FONT)
        f.set(qn("w:hAnsi"), FA_FONT)
        f.set(qn("w:cs"), FA_FONT)
        rp.insert(0, f)
        rp.append(_el("w:szCs", val=str(int(size * 2))))
        rp.append(OxmlElement("w:bCs"))


INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
LRM = "\u200e"
DIG = r"0-9\u06F0-\u06F9"
NUM_RUN = re.compile(
    r"[%s]+(?:\s?[.,\u066B\u066C:/\-\u2013\u2014\u00B1~\u00D7=+%%]\s?[%s]+)+" % (DIG, DIG))
# The lookbehind excludes Latin letters and LRM as well as digits: without that,
# the "-50" inside "ResNet-50" is treated as a signed number and gets its own
# LRM pair, splitting the model name into "ResNet" + "-50".
LEAD_SIGN = re.compile(
    r"(?<![%sA-Za-z‎])([±~<>≥≤≈+−-]\s?[%s]+"
    r"(?:[.,٫][%s]+)?)" % (DIG, DIG, DIG))
LATIN_PAREN = re.compile(r"\(([A-Za-z][A-Za-z0-9 .,\-/&'\u2013]*)\)")
# Latin identifiers carrying digits: ResNet-50, ViT-B/16, DINOv2-S. NUM_RUN
# cannot catch these because it needs digits on BOTH sides of a separator, so
# the digits stay inside the right-to-left run and Word renders them with
# Arabic-Indic glyphs - "ResNet-50" comes out as "ResNet-\u0665\u0660".
LATIN_ID = re.compile(r"(?:[A-Za-z][A-Za-z0-9]*(?:[-/.][A-Za-z0-9]+)+|[A-Za-z]+[0-9][A-Za-z0-9]*|[0-9]+[A-Za-z][A-Za-z0-9]*(?:\.[0-9]+)?)")
# capturing form, so re.split keeps the identifier as its own segment
LATIN_SEG = re.compile(r"([A-Za-z][A-Za-z0-9]*(?:[-/.][A-Za-z0-9]+)+|[A-Za-z]+[0-9][A-Za-z0-9]*|[0-9]+[A-Za-z][A-Za-z0-9]*(?:\.[0-9]+)?)")


def _fix_plain(t):
    t = NUM_RUN.sub(lambda m: LRM + m.group(0) + LRM, t)
    t = LEAD_SIGN.sub(lambda m: LRM + m.group(1) + LRM, t)
    t = LATIN_PAREN.sub(lambda m: LRM + "(" + m.group(1) + ")" + LRM, t)
    return t


def fix_bidi(t):
    # Latin identifiers (E67, 22q11.2, ResNet-50) are kept whole: the number
    # rules must not insert marks inside them, or they split into two runs.
    out = []
    for seg in LATIN_SEG.split(t):
        if not seg:
            continue
        if LATIN_SEG.fullmatch(seg):
            out.append(LRM + seg + LRM if any(c.isdigit() for c in seg) else seg)
        else:
            out.append(_fix_plain(seg))
    return "".join(out)


def add_inline(p, text, size=10.5, base_bold=False, base_italic=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            part, bold, italic = part[3:-3], True, True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            inner = part[2:-2]
            if "*" in inner:
                add_inline(p, inner, size, True, base_italic, color)
                continue
            part, bold = inner, True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            part = part[1:part.index("](")]
        if not mono:
            part = fix_bidi(part)
        # Word substitutes Arabic-Indic glyphs for ASCII digits inside any run
        # marked w:rtl, so "ResNet-50" renders as "ResNet-٥٠". LRM does not help
        # because this is numeral SHAPING, not bidi ordering. Latin identifiers
        # carrying digits are therefore emitted as their own runs with w:rtl
        # explicitly cleared - the document default sets it, so it must be
        # overridden rather than merely omitted.
        segs = [part] if mono else [s for s in LATIN_SEG.split(part) if s]
        for seg in segs:
            r = p.add_run(seg)
            if mono:
                run_font(r, MONO_FONT, MONO_FONT, size - 1)
            elif LATIN_SEG.fullmatch(seg) and any(c.isdigit() for c in seg):
                run_font(r, size=size)
                r._element.get_or_add_rPr().append(_el("w:rtl", val="0"))
            else:
                run_font(r, size=size)
                run_rtl(r)
            if bold:
                bold_cs(r)
            if italic:
                r.italic = True
                r._element.get_or_add_rPr().append(OxmlElement("w:iCs"))
            if color is not None:
                r.font.color.rgb = color


def add_inline_ltr(p, text, size=10):
    """Left-to-right runs for Latin-only paragraphs (reference entries)."""
    for part in INLINE.split(text):
        if not part:
            continue
        italic = bold = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        r = p.add_run(part)
        run_font(r, size=size)
        r._element.get_or_add_rPr().append(_el("w:rtl", val="0"))
        if bold:
            bold_cs(r)
        if italic:
            r.italic = True


def split_row(line):
    line = line.strip().strip("|")
    return [c.strip() for c in line.split("|")]


def is_sep(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|[\s:|-]*$", line)) and "-" in line


def build_table(doc, rows):
    hdr, body = rows[0], rows[1:]
    n = len(hdr)
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_rtl(t)
    for i, txt in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        para_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, txt, 9.5, True)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A5F")
        cell_borders(c, "1F3A5F")
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(n):
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(2)
            txt = row[i] if i < len(row) else ""
            if txt and not re.search(r"[؀-ۿ]", txt):
                # Latin-only cell: lay it out left-to-right
                p._p.get_or_add_pPr().append(_el("w:bidi", val="0"))
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_ltr(p, txt, 9)
            else:
                add_inline(p, txt, 9.5)
            if ri % 2 == 1:
                shade(c, "F2F5F9")
            cell_borders(c)


def convert(md_path, out_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    set_doc_defaults(doc)
    add_page_numbers(doc)
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.0)

    i, in_code, buf = 0, False, []
    while i < len(lines):
        line = lines[i]
        st = line.strip()

        if st.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para_shading(p, "F4F6F8")
                for k, cl in enumerate(buf):
                    r = p.add_run(cl)
                    run_font(r, MONO_FONT, MONO_FONT, 8.5)
                    if k < len(buf) - 1:
                        r.add_break()
                buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            buf.append(line)
            i += 1
            continue
        if not st:
            i += 1
            continue

        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            para_border_bottom(p, "C9D2DD", "6")
            i += 1
            continue

        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", st)
        if m:
            alt, rel = m.group(1), m.group(2)
            path = rel if os.path.isabs(rel) else os.path.normpath(
                os.path.join(os.path.dirname(md_path), rel))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            try:
                p.add_run().add_picture(path, width=Inches(6.2))
            except Exception:
                r = p.add_run("[image missing: %s]" % rel)
                run_font(r, size=9)
            if alt:
                cap = doc.add_paragraph()
                para_rtl(cap)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(12)
                add_inline(cap, alt, 9, True, color=MUTED)
            i += 1
            continue

        if st.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            rows = [split_row(st)]
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(split_row(lines[j]))
                j += 1
            build_table(doc, rows)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i = j
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", st)
        if m:
            lvl = min(len(m.group(1)), 4)
            p = doc.add_paragraph(style="Heading %d" % lvl)
            para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            keep_with_next(p)
            p.paragraph_format.space_before = Pt(15 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, m.group(2), {1: 16, 2: 13, 3: 11.5, 4: 11}[lvl], True, color=ACCENT)
            if lvl <= 2:
                para_border_bottom(p)
            i += 1
            continue

        if st.startswith(">"):
            b = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                b.append(lines[i].strip().lstrip(">").strip())
                i += 1
            # an empty ">" line separates paragraphs inside one box
            groups, cur = [], []
            for x in b + [""]:
                if x:
                    cur.append(x)
                elif cur:
                    groups.append(" ".join(cur))
                    cur = []
            for text in groups:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.right_indent = Cm(0.7)
                p.paragraph_format.space_after = Pt(2 if len(groups) > 1 else 6)
                para_shading(p, "EEF3F8")
                if not re.search(r"[؀-ۿ]", text):
                    # Latin-only paragraph: lay it out left-to-right so brackets
                    # and operators are not reordered by the bidi algorithm
                    p._p.get_or_add_pPr().append(_el("w:bidi", val="0"))
                    formula = bool(re.search(r"[=√²⁻]", text))
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if formula else WD_ALIGN_PARAGRAPH.LEFT
                    plain = re.sub(r"\*+", "", text)
                    font = "Cambria Math" if formula else FA_FONT
                    r = p.add_run(plain)
                    run_font(r, font, font, 11 if formula else 10)
                    r._element.get_or_add_rPr().append(_el("w:rtl", val="0"))
                    if text.startswith("**"):
                        bold_cs(r)
                    if text.startswith("*") and not text.startswith("**"):
                        r.italic = True
                    continue
                para_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_inline(p, text, 10)
            continue

        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(2))
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m and not re.search(r"[؀-ۿ]", m.group(3)):
            p = doc.add_paragraph()
            p._p.get_or_add_pPr().append(_el("w:bidi", val="0"))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(3)
            add_inline_ltr(p, m.group(2) + ". " + m.group(3), 9.5)
            i += 1
            continue
        if m:
            p = doc.add_paragraph()
            para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(3)
            lead = p.add_run(m.group(2) + ". ")
            run_font(lead, size=10.5)
            run_rtl(lead)
            bold_cs(lead)
            add_inline(p, m.group(3))
            i += 1
            continue

        p = doc.add_paragraph()
        para_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0.6)
        add_inline(p, st)
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    convert(SRC, OUT)
